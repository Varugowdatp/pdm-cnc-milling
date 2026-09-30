"""
Generates categorised Excel files for batch-upload testing, then runs every
file through the real POST /api/predict/batch endpoint (temp database) and
writes 00_TEST_INDEX.xlsx with what the app actually returned.

Usage (from the project root):
    .venv/Scripts/python.exe generate_batch_tests.py

Deletes and recreates batch_testing/ from scratch. Seeded, so the output is
identical on every run for the same model and dataset.
"""
from __future__ import annotations

import io
import math
import shutil
import sys
import tempfile
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
OUT = ROOT / "batch_testing"

from src.config import EXPANDED_CSV, RAW_CSV, THRESHOLDS  # noqa: E402

rng = np.random.default_rng(2026)
CANON = ["machine_id", "machine_type", "air_temperature", "process_temperature",
         "rotational_speed", "torque", "tool_wear"]
REGISTRY: list[dict] = []


# ------------------------------------------------------------------ helpers
def ramp(a, b, n):
    return np.linspace(a, b, n)


def noise(n, sd):
    return rng.normal(0, sd, n)


def run(mid, mtype, n, air, delta, speed, torque, wear, sd=(0.2, 0.1, 8, 0.8, 0)):
    """Build one machine's run. Each sensor arg is a scalar or (start, end) ramp."""
    def series(v, s):
        base = ramp(*v, n) if isinstance(v, tuple) else np.full(n, float(v))
        return base + (noise(n, s) if s else 0)
    a = series(air, sd[0])
    p = a + series(delta, sd[1])
    return pd.DataFrame({
        "machine_id": mid,
        "machine_type": mtype,
        "air_temperature": np.round(a, 1),
        "process_temperature": np.round(p, 1),
        "rotational_speed": np.round(series(speed, sd[2])).astype(int),
        "torque": np.round(np.clip(series(torque, sd[3]), 3, 80), 1),
        "tool_wear": np.round(np.clip(series(wear, sd[4]), 0, 260)).astype(int),
    })


def save(rel, df, purpose, expect, sheet="Readings", status=200, raw_bytes=None):
    path = OUT / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    if raw_bytes is not None:
        path.write_bytes(raw_bytes)
    else:
        with pd.ExcelWriter(path, engine="openpyxl") as xw:
            df.to_excel(xw, sheet_name=sheet, index=False)
        style_sheet(path)
    REGISTRY.append({"file": rel, "purpose": purpose, "expect": expect,
                     "expect_status": status, "rows": 0 if df is None else len(df)})


def style_sheet(path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if ws.max_row > 25000:
            continue
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(horizontal="center")
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(col[0].value or "")), 10) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
        ws.freeze_panes = "A2"
    wb.save(path)


# ------------------------------------------------------------------ 01 healthy
def healthy():
    base = "01_Healthy_Machines"
    for t, name in [("L", "Type_L_Low_Quality"), ("M", "Type_M_Medium_Quality"),
                    ("H", "Type_H_High_Quality")]:
        df = run(f"{t}-HLT-01", t, 60, 298.5, 10.2, 1520, 40, (5, 90))
        save(f"{base}/{name}/{t}_single_machine_60_readings.xlsx", df,
             f"One healthy type-{t} machine, 60 readings, all sensors in the safe envelope, wear 5->90 min.",
             "All rows Normal, alarm NORMAL, high health.")
        fleet = pd.concat([run(f"{t}-FLT-{i:02d}", t, 25, 297 + i * 0.4, 10 + rng.uniform(-0.5, 1.0),
                               1450 + i * 25, 36 + i, (10, 80)) for i in range(1, 6)])
        save(f"{base}/{name}/{t}_fleet_5_machines.xlsx", fleet,
             f"Five healthy type-{t} machines at different stable operating points, 25 readings each.",
             "All rows Normal; summary reports 5 machines.")
    snap = pd.DataFrame([{"machine_id": f"SNAP-{i:03d}", "machine_type": "LMH"[i % 3],
                          "air_temperature": round(rng.uniform(296.5, 301), 1),
                          "process_temperature": 0, "rotational_speed": int(rng.uniform(1420, 1650)),
                          "torque": round(rng.uniform(32, 46), 1), "tool_wear": int(rng.uniform(0, 120))}
                         for i in range(100)])
    snap["process_temperature"] = (snap["air_temperature"] + rng.uniform(9.8, 11.5, 100)).round(1)
    save(f"{base}/Mixed_Types/healthy_snapshot_100_machines.xlsx", snap,
         "100 unrelated healthy machines, one reading each, mixed L/M/H.",
         "All or almost all rows Normal.")


# ------------------------------------------------------------------ 02 synthetic failure modes
def failure_modes():
    base = "02_Failure_Modes_Synthetic"
    # HDF: delta < 8.6 K AND speed < 1380 rpm
    d = f"{base}/HDF_Heat_Dissipation"
    save(f"{d}/HDF_early_warning_approach.xlsx",
         run("HDF-EW-01", "L", 40, 300.5, (10.4, 8.9), (1500, 1390), 44, (40, 60)),
         "Temp delta falls 10.4->8.9 K and speed 1500->1390 rpm: approaching but not crossing the HDF limit.",
         "Health declines through the run; later rows may raise an early HDF warning.")
    save(f"{d}/HDF_confirmed_failure.xlsx",
         run("HDF-FL-01", "M", 30, 301.5, (8.3, 7.6), (1360, 1320), 50, (60, 70)),
         "Temp delta 8.3->7.6 K with speed 1360->1320 rpm: inside the HDF zone on every row.",
         "Rows predicted HDF, alarm CRITICAL.")
    save(f"{d}/HDF_full_degradation_run.xlsx",
         run("HDF-RUN-01", "L", 120, (298, 301.8), (11.0, 7.8), (1560, 1330), (38, 48), (10, 70)),
         "Healthy start degrading all the way into HDF over 120 readings.",
         "Normal at the start, HDF by the end, alarms escalate.")

    # PWF: power = torque * speed * 2pi/60 outside 3500-9000 W
    d = f"{base}/PWF_Power_Failure"
    save(f"{d}/PWF_overload_high_power.xlsx",
         run("PWF-OV-01", "M", 30, 299, 10, (1480, 1420), (62, 66), (50, 60)),
         "Torque 62->66 Nm at ~1450 rpm: mechanical power above 9000 W.",
         "Rows predicted PWF, alarm CRITICAL.")
    save(f"{d}/PWF_underload_low_power.xlsx",
         run("PWF-UN-01", "L", 30, 299, 10.5, (2550, 2750), (12, 10), (30, 40), sd=(0.2, 0.1, 15, 0.3, 0)),
         "Torque 12->10 Nm at ~2650 rpm: mechanical power below 3500 W.",
         "Rows predicted PWF, alarm CRITICAL.")
    save(f"{d}/PWF_overload_degradation_run.xlsx",
         run("PWF-RUN-01", "H", 120, 298.8, 10.3, (1500, 1430), (40, 66), (10, 80)),
         "Torque climbs 40->66 Nm over 120 readings, pushing power past 9000 W.",
         "Normal at the start, PWF by the end.")

    # OSF: tool_wear * torque > {L:11000, M:12000, H:13000}
    d = f"{base}/OSF_Overstrain"
    for t, wear, tq in [("L", (150, 200), (54, 60)), ("M", (160, 205), (56, 63)),
                        ("H", (175, 215), (58, 64))]:
        lim = THRESHOLDS["osf_limit"][t]
        save(f"{d}/OSF_type_{t}_limit_{int(lim)}.xlsx",
             run(f"OSF-{t}-01", t, 50, 299.5, 10.3, (1340, 1300), tq, wear),
             f"Type {t}: tool wear x torque climbs past the {int(lim)} min-Nm overstrain limit (power kept under 9000 W).",
             "Early rows Normal; final rows predicted OSF with CRITICAL alarm and hard breaches.")
    save(f"{d}/OSF_same_readings_all_three_types.xlsx",
         pd.concat([run(f"OSF-CMP-{t}", t, 20, 299.5, 10.3, 1300, 60, (185, 210), sd=(0, 0, 0, 0, 0))
                    for t in "LMH"]),
         "Identical readings (torque 60, wear 185->210) on L, M and H machines - strain 11100->12600.",
         "Model flags OSF on all three, but hard breaches differ by type: L on every row, M only once "
         "strain passes 12000, H never (limit 13000) - shows the per-type limit.")

    # TWF: wear inside 200-240 min
    d = f"{base}/TWF_Tool_Wear"
    save(f"{d}/TWF_wear_ramp_0_to_250.xlsx",
         run("TWF-RUN-01", "M", 126, 299, 10.2, 1500, 38, (0, 250), sd=(0.2, 0.1, 8, 0.8, 0)),
         "Tool wear counts 0->250 min at a normal load, every other sensor healthy.",
         "Normal while wear is low; TWF / maintenance advisory as wear enters 200-240 min.")
    save(f"{d}/TWF_inside_wear_window.xlsx",
         run("TWF-WIN-01", "L", 41, 298.7, 10.1, 1510, 35, (200, 240)),
         "Every reading inside the 200-240 min tool-wear window at a light load.",
         "Mostly TWF predictions (tool replacement due).")
    after = run("TWF-CHG-01", "M", 60, 299, 10.2, 1500, 38, 0)
    after["tool_wear"] = np.r_[np.linspace(180, 236, 30), np.linspace(0, 29, 30)].round().astype(int)
    save(f"{d}/TWF_tool_change_reset.xlsx", after,
         "Wear climbs 180->236 min, then the tool is changed and the counter resets to 0.",
         "TWF risk near the end of the first half, back to Normal after the reset.")


# ------------------------------------------------------------------ 03 real trajectories
def real_trajectories(df):
    base = "03_Run_To_Failure_Real_Dataset"
    syn = df[df["source"] != "AI4I-2020"]
    first_fail = (syn[syn["failure_class"] != "Normal"].groupby("machine_id")["failure_class"].first())
    modes = syn.groupby("machine_id")["degradation_mode"].first()
    picks = {
        "HDF_Heat_Dissipation": [m for m in modes.index if modes[m] == "HDF" and first_fail.get(m) == "HDF"][:2],
        "PWF_Overload": [m for m in modes.index if modes[m] == "PWF:overload" and first_fail.get(m) == "PWF"][:2],
        "PWF_Underload": [m for m in modes.index if modes[m] == "PWF:underload" and first_fail.get(m) == "PWF"][:2],
        "OSF_Overstrain": [m for m in modes.index if modes[m] == "OSF" and first_fail.get(m) == "OSF"][:2],
        "TWF_Tool_Wear": [m for m in modes.index if modes[m] == "TWF" and first_fail.get(m) == "TWF"][:2],
    }
    cols = CANON + ["cycle", "failure_class", "label_horizon", "rul"]
    for folder, mids in picks.items():
        for mid in mids:
            r = syn[syn["machine_id"] == mid].sort_values("cycle")[cols]
            r = r.rename(columns={"cycle": "actual_cycle", "failure_class": "actual_failure_class",
                                  "label_horizon": "actual_label_horizon", "rul": "actual_rul"})
            fc = first_fail[mid]
            save(f"{base}/{folder}/{mid}_type_{r['machine_type'].iloc[0]}_{len(r)}_readings.xlsx", r,
                 f"Complete simulated life of {mid} ending in {fc}. 'actual_*' columns are ground truth (ignored by the app).",
                 f"Health falls over the run; {fc} predicted before/at the failure. Compare with actual_* columns.")
    all_modes = pd.concat([syn[syn["machine_id"] == ms[0]].sort_values("cycle")[CANON]
                           for ms in picks.values()])
    save(f"{base}/All_Modes_Combined/five_machines_one_per_mode.xlsx", all_modes,
         "One full run-to-failure machine per failure mode stacked in one file.",
         "Summary shows 5 machines with a mix of HDF/PWF/OSF/TWF predictions.")


# ------------------------------------------------------------------ 04 mixed fleet
def mixed_fleet():
    base = "04_Mixed_Fleet"
    parts = [
        run("LINE-A-01", "L", 40, 298.5, 10.2, 1520, 40, (5, 60)),
        run("LINE-A-02", "M", 40, 299.0, 10.1, 1490, 42, (20, 80)),
        run("LINE-A-03", "H", 40, 298.8, 10.4, 1550, 38, (0, 50)),
        run("LINE-B-01", "L", 40, (300, 301.8), (9.8, 7.8), (1450, 1330), 46, (40, 80)),
        run("LINE-B-02", "M", 40, 299, 10, (1480, 1420), (48, 66), (30, 60)),
        run("LINE-B-03", "L", 40, 299.5, 10, (1420, 1360), (54, 60), (150, 200)),
        run("LINE-C-01", "H", 40, 299, 10.2, 1500, 38, (170, 245)),
        run("LINE-C-02", "M", 40, 298.2, 10.6, 1600, 36, (0, 40)),
    ]
    fleet = pd.concat(parts, ignore_index=True)
    save(f"{base}/plant_8_machines_grouped.xlsx", fleet,
         "8 machines on 3 lines, rows grouped by machine: 4 healthy, 1 HDF, 1 PWF, 1 OSF, 1 TWF trend.",
         "Summary: 8 machines, mix of Normal and failure classes; healthy lines stay NORMAL.")
    inter = pd.concat([p.assign(_k=np.arange(len(p))) for p in parts]).sort_values(["_k", "machine_id"])
    save(f"{base}/plant_8_machines_interleaved.xlsx", inter.drop(columns="_k").reset_index(drop=True),
         "Same 8 machines but rows interleaved in time order (as a data logger would export them).",
         "Per-machine results identical to the grouped file - grouping by machine_id is order-independent.")
    shift = pd.concat([p.iloc[:12] for p in parts], ignore_index=True)
    save(f"{base}/shift_report_12_readings_each.xlsx", shift,
         "Only the first 12 readings (one 2-hour shift) of each of the 8 machines.",
         "Degrading machines are still early; mostly Normal/ADVISORY.")


# ------------------------------------------------------------------ 05 real AI4I
def real_ai4i(raw):
    base = "05_Real_AI4I_2020_Data"
    # The app aliases BOTH 'UDI' and 'Product ID' to machine_id, which creates a
    # duplicate column and rejects the file. Keep Product ID only here; the
    # full-header file lives in 10_Known_App_Issues.
    full_raw = raw
    raw = raw.drop(columns="UDI")
    save(f"{base}/Original_UCI_Columns/ai4i_first_500_rows.xlsx", raw.head(500),
         "First 500 rows of the real UCI AI4I 2020 file with its original column names (Product ID, Type, "
         "Air temperature [K] ...). The UDI column is removed - see 10_Known_App_Issues.",
         "Accepted via column aliases; mostly Normal predictions.")
    save(f"{base}/Original_UCI_Columns/ai4i_all_339_failures.xlsx", raw[raw["Machine failure"] == 1],
         "All 339 real rows labelled 'Machine failure = 1', original column names and failure flags.",
         "Large share predicted as a failure class; compare with TWF/HDF/PWF/OSF flag columns.")
    for flag in ["HDF", "PWF", "OSF", "TWF"]:
        sub = raw[raw[flag] == 1]
        save(f"{base}/By_Failure_Flag/ai4i_{flag}_rows_{len(sub)}.xlsx", sub,
             f"The {len(sub)} real AI4I rows flagged {flag}.",
             f"Most rows predicted {flag} (TWF is stochastic in AI4I, so lower agreement there is expected).")
    normal = raw[raw["Machine failure"] == 0].sample(500, random_state=7)
    save(f"{base}/By_Failure_Flag/ai4i_normal_random_500.xlsx", normal,
         "500 random real rows with no failure flag.",
         "Nearly all rows Normal.")
    for t in "LMH":
        sub = raw[raw["Type"] == t].sample(300, random_state=11)
        save(f"{base}/By_Machine_Type/ai4i_type_{t}_random_300.xlsx", sub,
             f"300 random real rows of product type {t}.", "Mostly Normal; a few failures where flagged.")

    save("10_Known_App_Issues/ai4i_full_original_header_UDI_and_Product_ID.xlsx", full_raw.head(200),
         "APP ISSUE: the unmodified UCI header contains both 'UDI' and 'Product ID'. The app aliases both to "
         "machine_id, producing two machine_id columns.",
         "Should be accepted. CURRENTLY rejected with HTTP 422 'Grouper for machine_id not 1-dimensional'. "
         "Workaround: delete the UDI column.", status=422)


# ------------------------------------------------------------------ 06 boundaries
def boundaries():
    base = "06_Boundary_Edge_Cases"
    def row(mid, t, air, proc, spd, tq, wear, note):
        return {"machine_id": mid, "machine_type": t, "air_temperature": air, "process_temperature": proc,
                "rotational_speed": spd, "torque": tq, "tool_wear": wear, "note": note}

    hdf = pd.DataFrame([
        row("B-HDF-1", "L", 300.0, 308.7, 1379, 45, 50, "delta 8.7 K (just safe), speed 1379"),
        row("B-HDF-2", "L", 300.0, 308.6, 1379, 45, 50, "delta exactly 8.6 K, speed 1379"),
        row("B-HDF-3", "L", 300.0, 308.5, 1379, 45, 50, "delta 8.5 K (breach), speed 1379"),
        row("B-HDF-4", "L", 300.0, 308.5, 1380, 45, 50, "delta 8.5 K, speed exactly 1380"),
        row("B-HDF-5", "L", 300.0, 308.5, 1381, 45, 50, "delta 8.5 K, speed 1381 (safe)"),
    ])
    save(f"{base}/HDF_threshold_8.6K_and_1380rpm.xlsx", hdf,
         "Readings placed exactly on and either side of the HDF limits (delta 8.6 K AND speed 1380 rpm). Extra 'note' column is ignored.",
         "Only rows with BOTH delta < 8.6 and speed < 1380 are hard HDF breaches.")

    k = 60 / (2 * math.pi)
    pw = []
    for watts, note in [(3499, "3499 W (breach)"), (3499.5, "3499.5 W (breach)"), (3500.5, "3500.5 W (safe)"),
                        (3501, "3501 W (safe)"), (8999, "8999 W (safe)"), (8999.5, "8999.5 W (safe)"),
                        (9000.5, "9000.5 W (breach)"), (9001, "9001 W (breach)")]:
        tq = round(watts * k / 1500, 7)
        pw.append(row(f"B-PWF-{watts}", "M", 299.0, 309.5, 1500, tq, 40, note))
    save(f"{base}/PWF_threshold_3500W_and_9000W.xlsx", pd.DataFrame(pw),
         "Torque set at 1500 rpm so mechanical power lands 0.5 W and 1 W either side of the 3500 W and 9000 W "
         "limits (exactly 3500.000 W cannot be written as torque x rpm).",
         "Model may still say Normal (1 W is invisible to it), but exactly 4 rows - 3499, 3499.5, 9000.5, "
         "9001 W - are hard power breaches with a CRITICAL alarm.")

    osf = []
    for t in "LMH":
        lim = THRESHOLDS["osf_limit"][t]
        for strain, tag in [(lim - 100, "below"), (lim, "exact"), (lim + 100, "above")]:
            wear = 200
            osf.append(row(f"B-OSF-{t}-{tag}", t, 299.0, 309.5, 1300, round(strain / wear, 2), wear,
                           f"type {t}: strain {int(strain)} vs limit {int(lim)} ({tag})"))
    save(f"{base}/OSF_threshold_per_type.xlsx", pd.DataFrame(osf),
         "Tool wear 200 min at 1300 rpm (power stays under 9000 W) with torque chosen so strain sits 100 "
         "below / exactly on / 100 above each type's limit.",
         "Model predicts OSF on all rows (they are all near the limit); only the 'above' row per type is a "
         "hard overstrain breach.")

    twf = pd.DataFrame([row(f"B-TWF-{w}", "M", 299.0, 309.3, 1500, 38, w, f"tool wear {w} min")
                        for w in [0, 199, 200, 201, 220, 239, 240, 241, 253]])
    save(f"{base}/TWF_wear_window_200_240.xlsx", twf,
         "Tool wear at 0, 199, 200, 201, 220, 239, 240, 241, 253 min, everything else nominal.",
         "Risk rises inside/after the 200-240 window; wear 0 and 199 read Normal.")

    rng_rows = pd.DataFrame([
        row("B-RNG-MIN", "L", 290.0, 300.0, 1000, 3.0, 0, "every sensor at its HMI minimum"),
        row("B-RNG-MAX", "H", 315.0, 320.0, 3000, 80.0, 260, "every sensor at its HMI maximum"),
        row("B-RNG-OUT", "M", 330.0, 345.0, 3500, 95.0, 300, "outside the HMI range (still numeric)"),
        row("B-RNG-DEC", "M", 298.123456, 308.654321, 1551.7, 42.85, 108.5, "many decimal places"),
    ])
    save(f"{base}/sensor_range_extremes.xlsx", rng_rows,
         "Sensor values at the HMI minimum, maximum, beyond the range, and with long decimals.",
         "All 4 rows scored (no range rejection); extreme rows get failure/breach results.")


# ------------------------------------------------------------------ 07 column formats
def column_formats():
    base = "07_Column_Format_Variants"
    df = run("FMT-01", "M", 20, 299, 10.2, 1510, 40, (10, 40))
    save(f"{base}/canonical_column_names.xlsx", df,
         "Exact column names from the app's template.", "20 rows scored.")
    ai = df.rename(columns={"machine_id": "Product ID", "machine_type": "Type",
                            "air_temperature": "Air temperature [K]", "process_temperature": "Process temperature [K]",
                            "rotational_speed": "Rotational speed [rpm]", "torque": "Torque [Nm]",
                            "tool_wear": "Tool wear [min]"})
    save(f"{base}/uci_ai4i_column_names.xlsx", ai,
         "UCI AI4I header names (Product ID, Type, Air temperature [K] ...).", "20 rows scored via aliases.")
    sp = df.rename(columns={"machine_id": "Machine", "machine_type": "Variant", "air_temperature": "air_temp",
                            "process_temperature": "process_temp", "rotational_speed": "RPM",
                            "torque": "Torque", "tool_wear": "ToolWear"})
    save(f"{base}/spreadsheet_short_names.xlsx", sp,
         "Informal spreadsheet headers (Machine, Variant, air_temp, process_temp, RPM, Torque, ToolWear).",
         "20 rows scored via aliases.")
    sp2 = sp.rename(columns={"air_temp": "Air Temp", "process_temp": "Process Temp"})
    save("10_Known_App_Issues/short_names_with_spaces_Air_Temp.xlsx", sp2,
         "APP ISSUE: headers 'Air Temp' / 'Process Temp' (with a space). The alias table has 'air_temp' but the "
         "lookup runs before spaces are converted to underscores.",
         "Should be accepted. CURRENTLY rejected with HTTP 422 missing air_temperature, process_temperature. "
         "Workaround: write air_temp / process_temp.", status=422)
    up = df.rename(columns={c: c.upper().replace("_", " ") for c in df.columns})
    save(f"{base}/uppercase_with_spaces.xlsx", up,
         "Headers in UPPER CASE with spaces (AIR TEMPERATURE, TOOL WEAR ...).", "20 rows scored.")
    save(f"{base}/columns_reordered.xlsx", df[list(reversed(df.columns))],
         "Canonical names in reverse column order.", "20 rows scored; order does not matter.")
    save(f"{base}/no_machine_id_column.xlsx", df.drop(columns="machine_id"),
         "machine_id column omitted (it is optional).", "20 rows scored; summary reports 0 machines.")
    extra = df.assign(operator="Ravi", shift="A", notes="routine", humidity_pct=55)
    save(f"{base}/extra_unused_columns.xlsx", extra,
         "Valid data plus unrelated columns (operator, shift, notes, humidity_pct).", "20 rows scored; extras ignored.")
    lc = df.copy()
    lc["machine_type"] = [" m ", "m", "M", "l", "h"] * 4
    save(f"{base}/machine_type_lowercase_and_spaces.xlsx", lc,
         "machine_type written as 'm', ' m ', 'l', 'h'.", "Normalised to M/L/H; 20 rows scored.")
    bad_t = df.copy()
    bad_t["machine_type"] = ["X", "Medium", "", "Q"] * 5
    save(f"{base}/machine_type_unknown_values.xlsx", bad_t,
         "machine_type values the app does not know (X, Medium, blank, Q).",
         "Rows still scored; unknown types fall back to L.")
    save(f"{base}/custom_sheet_name.xlsx", df,
         "Data on a single sheet named 'Plant Log March' instead of Sheet1.",
         "20 rows scored (the first sheet is read regardless of name).", sheet="Plant Log March")


# ------------------------------------------------------------------ 08 invalid
def invalid():
    base = "08_Invalid_Inputs"
    good = run("BAD-01", "M", 10, 299, 10.2, 1510, 40, (10, 20))

    d = f"{base}/Expect_Rejected_File"
    save(f"{d}/missing_torque_column.xlsx", good.drop(columns="torque"),
         "Required 'torque' column removed.", "HTTP 422 - missing required column(s): torque.", status=422)
    save(f"{d}/missing_machine_type_column.xlsx", good.drop(columns="machine_type"),
         "Required 'machine_type' column removed.", "HTTP 422 - missing required column(s): machine_type.", status=422)
    save(f"{d}/missing_all_sensor_columns.xlsx", good[["machine_id", "machine_type"]],
         "Only machine_id and machine_type present.", "HTTP 422 listing all five sensor columns.", status=422)
    save(f"{d}/headers_only_no_rows.xlsx", good.iloc[0:0],
         "Correct headers but zero data rows.", "HTTP 422 - No rows found in the file.", status=422)
    allbad = good.copy().astype(object)
    allbad["torque"] = "high"
    save(f"{d}/all_rows_non_numeric.xlsx", allbad,
         "Every torque cell is the text 'high'.", "HTTP 422 - No valid rows found.", status=422)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"Info": ["Plant sensor export", "Data is on the next sheet"]}).to_excel(xw, sheet_name="Cover", index=False)
        good.to_excel(xw, sheet_name="Readings", index=False)
    save(f"{d}/data_on_second_sheet.xlsx", good,
         "First sheet is a cover page; the readings are on sheet 2.",
         "HTTP 422 - only the first sheet is read, so required columns are missing.", status=422,
         raw_bytes=buf.getvalue())
    save(f"{d}/corrupted_not_really_excel.xlsx", None,
         "A text file renamed to .xlsx.", "HTTP 422 - Could not parse the file.", status=422,
         raw_bytes=b"this is not an excel workbook,just,text\n1,2,3\n")

    d = f"{base}/Expect_Partial_Rows_Rejected"
    mix = good.copy().astype(object)
    mix.loc[2, "torque"] = "N/A"
    mix.loc[5, "air_temperature"] = "error"
    mix.loc[7, "tool_wear"] = None
    save(f"{d}/some_text_and_blank_cells.xlsx", mix,
         "10 rows; row 4 torque 'N/A', row 7 air temp 'error', row 9 tool wear blank (Excel row numbers).",
         "HTTP 200 - 7 rows scored, 3 rejected with their row numbers.")
    zero = good.copy()
    zero.loc[1, "rotational_speed"] = 0
    zero.loc[4, "rotational_speed"] = -1500
    save(f"{d}/zero_and_negative_speed.xlsx", zero,
         "Rows with rotational_speed 0 and -1500 rpm.", "HTTP 200 - 8 scored, 2 rejected as non-physical.")
    blank = pd.concat([good.iloc[:5], pd.DataFrame([{c: None for c in good.columns}] * 2), good.iloc[5:]],
                      ignore_index=True)
    save(f"{d}/empty_rows_in_middle.xlsx", blank,
         "Two completely empty rows between valid readings.", "HTTP 200 - 10 scored, 2 rejected.")


# ------------------------------------------------------------------ 09 volume
def volume(df):
    base = "09_Volume_Performance"
    syn = df[df["source"] != "AI4I-2020"][CANON]
    for n in [100, 1000, 5000]:
        save(f"{base}/rows_{n}.xlsx", syn.head(n),
             f"First {n} readings of the run-to-failure dataset (several complete machines).",
             f"HTTP 200 - {n} rows scored.")
    save(f"{base}/rows_20000_at_limit.xlsx", syn.head(20000),
         "Exactly 20,000 rows - the maximum the upload accepts.", "HTTP 200 - 20,000 rows scored (slow: tens of seconds).")
    save(f"{base}/rows_20001_over_limit.xlsx", syn.head(20001),
         "20,001 rows - one more than the limit.", "HTTP 413 - the limit is 20,000.", status=413)


# ------------------------------------------------------------------ verify + index
def verify():
    from fastapi.testclient import TestClient
    from backend.core import database as db
    tmp = Path(tempfile.mkdtemp()) / "batch_verify.sqlite3"
    db.DB_PATH = tmp
    db.init_db(tmp, force=True)
    from backend.app import app

    with TestClient(app) as c:
        for i, r in enumerate(REGISTRY, 1):
            p = OUT / r["file"]
            with open(p, "rb") as fh:
                resp = c.post("/api/predict/batch", files={"file": (p.name, fh.read(),
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            r["status"] = resp.status_code
            r["size_kb"] = round(p.stat().st_size / 1024, 1)
            body = resp.json()
            if resp.status_code == 200:
                s = body["summary"]
                r["scored"] = body["rows_scored"]
                r["rejected"] = body["rows_rejected"]
                r["machines"] = s["machines"]
                r["classes"] = ", ".join(f"{k} {v}" for k, v in s["by_class"].items() if v)
                r["alarms"] = ", ".join(f"{k} {v}" for k, v in s["by_alarm_level"].items() if v)
                r["mean_health"] = s["mean_health"]
                r["hard_breaches"] = sum(1 for x in body["results"] if x["has_hard_breach"])
                last = Counter(x["predicted_class"] for x in body["results"][-10:])
                r["last_rows"] = ", ".join(f"{k} {v}" for k, v in last.most_common())
                r["message"] = ""
            else:
                r["message"] = str(body.get("detail"))[:200]
            r["pass"] = "PASS" if r["status"] == r["expect_status"] else "CHECK"
            if r["file"].startswith("10_Known_App_Issues") and r["pass"] == "PASS":
                r["pass"] = "KNOWN ISSUE"
            print(f"[{i:02d}/{len(REGISTRY)}] {r['status']} {r['pass']} {r['file']}  "
                  f"{r.get('classes', r['message'])[:90]}")


def write_index():
    cols = [("No", 5), ("Category", 26), ("Sub-folder", 28), ("File", 44), ("Rows", 8), ("Size KB", 9),
            ("What it tests", 60), ("Expected behaviour", 50), ("Expected HTTP", 9), ("Actual HTTP", 9),
            ("Status check", 9), ("Rows scored", 9), ("Rows rejected", 9), ("Machines", 9),
            ("Predicted classes", 34), ("Alarm levels", 38), ("Hard breach rows", 9), ("Last 10 rows predicted", 26),
            ("Mean health", 9), ("Error message", 50)]
    rows = []
    for i, r in enumerate(REGISTRY, 1):
        parts = r["file"].split("/")
        rows.append([i, parts[0], "/".join(parts[1:-1]) or "-", parts[-1], r["rows"], r.get("size_kb"),
                     r["purpose"], r["expect"], r["expect_status"], r.get("status"), r.get("pass"),
                     r.get("scored"), r.get("rejected"), r.get("machines"), r.get("classes"),
                     r.get("alarms"), r.get("hard_breaches"), r.get("last_rows"), r.get("mean_health"), r.get("message")])
    idx = pd.DataFrame(rows, columns=[c for c, _ in cols])

    cats = idx.groupby("Category").agg(Files=("File", "count"), Rows=("Rows", "sum"),
                                       Status_OK=("Status check", lambda s: int((s == "PASS").sum())))
    cats = cats.reset_index()
    how = pd.DataFrame({"Step": [1, 2, 3, 4, 5, 6, 7], "How to use these files": [
        "Start the app:  .venv\\Scripts\\python.exe -m backend.app  ->  http://127.0.0.1:8000/",
        "Open the Batch Upload screen.",
        "Pick a file from a category folder below and upload it.",
        "Compare the results screen with the 'Expected behaviour' and 'Actual' columns on the Test Index sheet.",
        "'Actual' columns were recorded by uploading every file to the real /api/predict/batch endpoint on "
        "the current model (temporary database), so they are what a correct run looks like.",
        "Files in 08_Invalid_Inputs/Expect_Rejected_File are SUPPOSED to be rejected with an error message.",
        "Files in 10_Known_App_Issues reproduce two input-handling problems found while building these tests; "
        "they are rejected today but should be accepted. Their workaround is in the Expected behaviour column.",
    ]})

    path = OUT / "00_TEST_INDEX.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        how.to_excel(xw, sheet_name="Read Me", index=False)
        cats.to_excel(xw, sheet_name="Categories", index=False)
        idx.to_excel(xw, sheet_name="Test Index", index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
    wb["Read Me"].column_dimensions["B"].width = 120
    wb["Categories"].column_dimensions["A"].width = 36
    ws = wb["Test Index"]
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = ws.dimensions
    green, amber = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFEB9C")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[10].fill = green if row[10].value == "PASS" else amber
    wb.save(path)


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir()
    df = pd.read_csv(EXPANDED_CSV)
    raw = pd.read_csv(RAW_CSV)
    healthy(); failure_modes(); real_trajectories(df); mixed_fleet(); real_ai4i(raw)
    boundaries(); column_formats(); invalid(); volume(df)
    print(f"{len(REGISTRY)} files written; verifying through the API ...")
    verify()
    write_index()
    print("index written")


if __name__ == "__main__":
    main()
